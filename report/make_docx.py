"""Convert a (simple) Markdown file to .docx: headings, paragraphs, bullets, numbered lists, tables, code blocks.
    python report/make_docx.py GUIDE.md report/GUIDE.docx      (needs: pip install python-docx)
"""
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches

import os
src, out = sys.argv[1], sys.argv[2]
COVER = "--cover" in sys.argv
base_dir = os.path.dirname(os.path.abspath(src))
lines = open(src, encoding="utf-8").read().split("\n")
doc = Document()
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Inches(0.9)
sec.top_margin = sec.bottom_margin = Inches(0.8)
doc.styles["Normal"].font.name = "Calibri"
doc.styles["Normal"].font.size = Pt(10.5)
for lvl, size in ((1, 20), (2, 15), (3, 12.5)):
    st = doc.styles[f"Heading {lvl}"]
    st.font.name = "Calibri"; st.font.size = Pt(size); st.font.color.rgb = RGBColor(0x2B, 0x2F, 0x7E)


if COVER:
    for _ in range(5):
        doc.add_paragraph()
    t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("Autonomous Cyber Defense Using Multi-Agent Reinforcement Learning"); r.bold = True; r.font.size = Pt(26)
    r.font.color.rgb = RGBColor(0x2B, 0x2F, 0x7E)
    t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("A simulation study with single-agent PPO, independent PPO and MAPPO, "
                  "including generalization tests, ablations and an offline CICIDS2017 replay"); r.font.size = Pt(13)
    doc.add_page_break()


def shade(cell, hex_):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hex_)
    tcPr.append(sh)


def runs(par, text, base_bold=False, size=None):
    # tokens: **bold**, `code`, *italic*
    for tok in re.split(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*\s][^*]*\*)", text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            r = par.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith("`") and tok.endswith("`"):
            r = par.add_run(tok[1:-1]); r.font.name = "Consolas"; r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0x8A, 0x1F, 0x5C)
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            r = par.add_run(tok[1:-1]); r.italic = True
        else:
            r = par.add_run(tok)
        if base_bold:
            r.bold = True
        if size:
            r.font.size = Pt(size)


i = 0
while i < len(lines):
    ln = lines[i]
    if ln.startswith("```"):
        buf = []
        i += 1
        while i < len(lines) and not lines[i].startswith("```"):
            buf.append(lines[i]); i += 1
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.2)
        r = p.add_run("\n".join(buf)); r.font.name = "Consolas"; r.font.size = Pt(9)
        pPr = p._p.get_or_add_pPr(); sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:fill"), "F0F0F4"); pPr.append(sh)
        i += 1
        continue
    if ln.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s\-:|]+\|$", lines[i + 1].strip()):
        rows = []
        while i < len(lines) and lines[i].startswith("|"):
            if not re.match(r"^\|[\s\-:|]+\|$", lines[i].strip()):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
            i += 1
        t = doc.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        for ri, row in enumerate(rows):
            for ci, txt in enumerate(row[: len(rows[0])]):
                c = t.cell(ri, ci); c.text = ""
                runs(c.paragraphs[0], txt, base_bold=(ri == 0), size=9)
                if ri == 0:
                    shade(c, "DDE1F4")
        doc.add_paragraph()
        continue
    im = re.match(r"^!\[(.*)\]\((.*)\)\s*$", ln)
    if im:
        path = os.path.normpath(os.path.join(base_dir, im.group(2)))
        if os.path.exists(path):
            pp = doc.add_paragraph(); pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pp.add_run().add_picture(path, width=Inches(6.2))
            cp = doc.add_paragraph(); cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            rr = cp.add_run(im.group(1)); rr.italic = True; rr.font.size = Pt(9)
        i += 1
        continue
    m = re.match(r"^(#{1,3})\s+(.*)", ln)
    if m:
        doc.add_heading(re.sub(r"[*`]", "", m.group(2)), level=len(m.group(1)))
    elif re.match(r"^\s*[*-]\s+", ln):
        indent = len(ln) - len(ln.lstrip())
        txt = re.sub(r"^\s*[*-]\s+", "", ln)
        while i + 1 < len(lines) and re.match(r"^\s{2,}\S", lines[i + 1]) and not re.match(r"^\s*([*-]|\d+\.)\s", lines[i + 1]):
            i += 1; txt += " " + lines[i].strip()
        p = doc.add_paragraph(style="List Bullet")
        if indent >= 2:
            p.paragraph_format.left_indent = Inches(0.55)
        runs(p, txt)
    elif re.match(r"^\s*\d+\.\s+", ln):
        txt = re.sub(r"^\s*\d+\.\s+", "", ln)
        while i + 1 < len(lines) and re.match(r"^\s{2,}\S", lines[i + 1]) and not re.match(r"^\s*([*-]|\d+\.)\s", lines[i + 1]):
            i += 1; txt += " " + lines[i].strip()
        p = doc.add_paragraph(style="List Number")
        runs(p, txt)
    elif ln.strip() in ("---", ""):
        pass
    elif ln.startswith(">"):
        p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.3); runs(p, ln.lstrip("> "))
    else:
        # continuation lines of a wrapped paragraph: join until blank/structure
        buf = [ln.strip()]
        while i + 1 < len(lines) and lines[i + 1].strip() and not re.match(r"^(#|\||```|\s*[*-]\s|\s*\d+\.\s|---)", lines[i + 1]):
            i += 1; buf.append(lines[i].strip())
        runs(doc.add_paragraph(), " ".join(buf))
    i += 1

doc.core_properties.title = "Autonomous Cyber Defense Using MARL: Complete Guide"
doc.save(out)
print("saved", out)
